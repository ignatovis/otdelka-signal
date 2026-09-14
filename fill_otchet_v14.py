"""
Fill all empty data cells in ОТЧЕТ sheet of v14 template.
Pattern taken from filled 8.1.1 and 8.2.1 К1 sections.
"""
import openpyxl
from openpyxl.styles import PatternFill, Alignment

FILE = '03_Финальная версия - Отделка для SIGNAL v14.xlsx'
wb = openpyxl.load_workbook(FILE)
ws = wb['ОТЧЕТ']

fill_data = PatternFill(start_color='FFF2F2F2', end_color='FFF2F2F2', fill_type='solid')
align_center = Alignment(horizontal='center', vertical='center')

FORMATS = {
    3: '0.00',       # C: ВСЕГО м²
    4: '#,##0.00',   # D: План на сегодня
    5: '#,##0.00',   # E: Факт на сегодня
    6: '#,##0.00',   # F: Дельта
    7: '0.0%',       # G: Факт %
}

def fmt(cell, col_num):
    cell.fill = fill_data
    cell.alignment = align_center
    cell.number_format = FORMATS[col_num]

def fill_row(row, plan_sheet, fact_sheet, data_col):
    c = ws.cell(row=row, column=3)
    c.value = f"=SUM('{plan_sheet}'!{data_col}5:{data_col}90)"
    fmt(c, 3)

    d = ws.cell(row=row, column=4)
    d.value = (f"=IFERROR(SUM('{plan_sheet}'!{data_col}5:"
               f"INDEX('{plan_sheet}'!{data_col}5:{data_col}90,"
               f"MATCH(TODAY(),'{plan_sheet}'!A5:A90,1))),0)")
    fmt(d, 4)

    e = ws.cell(row=row, column=5)
    e.value = (f"=IFERROR(SUM('{fact_sheet}'!{data_col}5:"
               f"INDEX('{fact_sheet}'!{data_col}5:{data_col}90,"
               f"MATCH(TODAY(),'{fact_sheet}'!A5:A90,1))),0)")
    fmt(e, 5)

    f_cell = ws.cell(row=row, column=6)
    f_cell.value = f"=E{row}-D{row}"
    fmt(f_cell, 6)

    g = ws.cell(row=row, column=7)
    g.value = f'=IFERROR(E{row}/C{row},"0")'
    fmt(g, 7)


def fill_section(code_row, plan_chern, fact_chern, plan_chist, fact_chist,
                 chern_cols, chist_cols, zapol_col=None):
    for i, col in enumerate(chern_cols):
        fill_row(code_row + 2 + i, plan_chern, fact_chern, col)
    for i, col in enumerate(chist_cols):
        fill_row(code_row + 6 + i, plan_chist, fact_chist, col)
    if zapol_col:
        fill_row(code_row + 9, plan_chist, fact_chist, zapol_col)


# Column mappings: underground
CHERN_PODZ = {
    '8.1.1': ('B','C','D'), '8.1.2': ('E','F','G'), '8.1.3': ('H','I','J'),
    '8.1.4': ('K','L','M'), '8.1.5': ('N','O','P'), '8.1.6': ('Q','R','S'),
}
CHIST_PODZ = {
    '8.1.1': ('B','C','D'), '8.1.2': ('G','H','I'), '8.1.3': ('L','M','N'),
    '8.1.4': ('Q','R','S'), '8.1.5': ('V','W','X'), '8.1.6': ('AA','AB','AC'),
}

# Column mappings: К1/К2 above-ground
CHERN_K = {
    '8.2.1': ('B','C','D'), '8.2.2': ('E','F','G'), '8.2.3': ('H','I','J'),
    '8.2.4': ('K','L','M'), '8.2.5': ('N','O','P'), '8.2.6': ('Q','R','S'),
    '8.2.7': ('T','U','V'),
}
CHIST_K = {
    '8.2.1': ('B','C','D'), '8.2.2': ('G','H','I'), '8.2.3': ('L','M','N'),
    '8.2.4': ('Q','R','S'), '8.2.5': ('V','W','X'), '8.2.6': ('AA','AB','AC'),
    '8.2.7': ('AF','AG','AH'),
}
ZAPOL_K = {
    '8.2.2': 'J', '8.2.3': 'O', '8.2.4': 'T',
    '8.2.5': 'Y', '8.2.6': 'AD', '8.2.7': 'AI',
}


# ===================== FILL UNDERGROUND 8.1.2-8.1.6 =====================
print("=== ПОДЗЕМНАЯ ЧАСТЬ ===")
underground = [
    ('8.1.2', 16), ('8.1.3', 26), ('8.1.4', 36),
    ('8.1.5', 46), ('8.1.6', 56),
]

filled = 0
for code, cr in underground:
    fill_section(cr,
                 'ПЛАН ЧЕРН ПОДЗЕМ', 'ФАКТ ЧЕРН ПОДЗЕМ',
                 'ПЛАН ЧИСТ ПОДЗЕМ', 'ФАКТ ЧИСТ ПОДЗЕМ',
                 CHERN_PODZ[code], CHIST_PODZ[code])
    filled += 6
    print(f"  {code} (row {cr}): 6 data rows filled")


# ===================== FILL К1 8.2.2-8.2.7 =====================
print("\n=== НАДЗЕМНАЯ К1 ===")
k1_codes = [
    ('8.2.2', 78), ('8.2.3', 89), ('8.2.4', 100),
    ('8.2.5', 111), ('8.2.6', 122), ('8.2.7', 133),
]

for code, cr in k1_codes:
    zc = ZAPOL_K.get(code)
    fill_section(cr,
                 'ПЛАН ЧЕРН К1', 'ФАКТ ЧЕРН К1',
                 'ПЛАН ЧИСТ К1', 'ФАКТ ЧИСТ К1',
                 CHERN_K[code], CHIST_K[code], zc)
    n = 7 if zc else 6
    filled += n
    print(f"  {code} К1 (row {cr}): {n} data rows filled")


# ===================== FILL К2 8.2.1-8.2.7 =====================
print("\n=== НАДЗЕМНАЯ К2 ===")
k2_codes = [
    ('8.2.1', 146), ('8.2.2', 156), ('8.2.3', 167),
    ('8.2.4', 178), ('8.2.5', 189), ('8.2.6', 200), ('8.2.7', 211),
]

for code, cr in k2_codes:
    zc = ZAPOL_K.get(code)
    fill_section(cr,
                 'ПЛАН ЧЕРН К2', 'ФАКТ ЧЕРН К2',
                 'ПЛАН ЧИСТ К2', 'ФАКТ ЧИСТ К2',
                 CHERN_K[code], CHIST_K[code], zc)
    n = 7 if zc else 6
    filled += n
    print(f"  {code} К2 (row {cr}): {n} data rows filled")


wb.save(FILE)
print(f"\nDone! {filled} data rows filled ({filled * 5} cells).")
