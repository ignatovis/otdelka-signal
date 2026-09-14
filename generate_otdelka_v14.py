"""
Генератор шаблона "Отделка для SIGNAL" v14.
Параметры: 5 корпусов, вес черн=0.6, вес чист=0.4, дата старта=2026-06-01, данные=0.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

# ─── CONFIG ───────────────────────────────────────────────────────────────────
NUM_KORPUS = 5
WEIGHT_CHERN = 0.6
WEIGHT_CHIST = 0.4
START_DATE = datetime(2026, 6, 1)
DATA_ROWS = 86  # rows 5..90
OUTPUT = r'\\mr.ru\Service\Personal\ignatov_i\Documents\CloudCode\Отделка SIGNAL\Отделка для SIGNAL (5 корпусов).xlsx'

UNDERGROUND_CODES = [
    ('8.1.1', 'Отделка паркинг и рампы'),
    ('8.1.2', 'Отделка эвакуац. лестн. клетки подзем.'),
    ('8.1.3', 'Лифт. холлы, тамбур-шлюзы подзем.'),
    ('8.1.4', 'Технич. помещения подзем.'),
    ('8.1.5', 'Прочие помещения подзем.'),
    ('8.1.6', 'Коммерч. помещения подзем.'),
]

ABOVE_GROUND_CODES = [
    ('8.2.1', 'Отделка лобби/гранд-лобби'),
    ('8.2.2', 'Отделка эвакуац. лестн. клетки надзем.'),
    ('8.2.3', 'Лифт. холлы, тамбур-шлюзы надзем.'),
    ('8.2.4', 'Технич. помещения надзем.'),
    ('8.2.5', 'Прочие помещения надзем.'),
    ('8.2.6', 'Коммерч. помещения надзем.'),
    ('8.2.7', 'Паркинг и рампы надзем.'),
]

# ─── STYLES ───────────────────────────────────────────────────────────────────
FILL_DARK_BLUE = PatternFill('solid', fgColor='FF2F5496')
FILL_BLUE = PatternFill('solid', fgColor='FF4472C4')
FILL_PURPLE = PatternFill('solid', fgColor='FF7030A0')
FILL_LIGHT_BLUE = PatternFill('solid', fgColor='FFB4C6E7')
FILL_LIGHTEST_BLUE = PatternFill('solid', fgColor='FFD9E2F3')
FILL_GRAY = PatternFill('solid', fgColor='FFF2F2F2')
FILL_RED_WARN = PatternFill('solid', fgColor='FFFFC7CE')
FILL_YELLOW = PatternFill('solid', fgColor='FFFFF2CC')
FILL_SIGNAL_LABEL = PatternFill('solid', fgColor='FFD6DCE4')
FILL_SIGNAL_DATE = PatternFill('solid', fgColor='FFDAEEF3')
FILL_SIGNAL_DATA_HDR = PatternFill('solid', fgColor='FFB4C6E7')
FILL_SIGNAL_FACT = PatternFill('solid', fgColor='FFEAF0FB')

FONT_WHITE_BOLD_14 = Font(bold=True, size=14, color='FFFFFFFF')
FONT_WHITE_BOLD = Font(bold=True, color='FFFFFFFF')
FONT_BOLD = Font(bold=True)
FONT_RED_BOLD = Font(bold=True, color='FFFF0000')
FONT_NORMAL = Font()

ALIGN_CENTER = Alignment(horizontal='center', vertical='center')
ALIGN_LEFT_CENTER = Alignment(horizontal='left', vertical='center')
ALIGN_CENTER_WRAP = Alignment(horizontal='center', vertical='center', wrap_text=True)

THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)


def col_letter(n):
    return get_column_letter(n)


# ─── ОТЧЕТ ROW LAYOUT CALCULATOR ─────────────────────────────────────────────
def calc_otchet_layout():
    layout = {}
    layout['title_row'] = 1
    layout['warn_row'] = 2
    layout['underground_section_row'] = 4
    layout['underground_header_row'] = 5

    row = 6
    layout['underground'] = []
    for i, (code, name) in enumerate(UNDERGROUND_CODES):
        code_row = row
        layout['underground'].append({
            'code': code, 'name': name, 'idx': i,
            'code_row': code_row,
            'chern_label': code_row + 1,
            'chern_data': [code_row + 2, code_row + 3, code_row + 4],
            'chist_label': code_row + 5,
            'chist_data': [code_row + 6, code_row + 7, code_row + 8],
            'blank': code_row + 9,
        })
        row = code_row + 10

    layout['korpus_sections'] = []
    for k in range(1, NUM_KORPUS + 1):
        section = {'k': k, 'section_row': row, 'header_row': row + 1, 'codes': []}
        row += 2
        for j, (code, name) in enumerate(ABOVE_GROUND_CODES):
            code_row = row
            has_zapol = (code != '8.2.1')
            chist_data = [code_row + 6, code_row + 7, code_row + 8]
            if has_zapol:
                zapol_row = code_row + 9
                blank = code_row + 10
            else:
                zapol_row = None
                blank = code_row + 9
            section['codes'].append({
                'code': code, 'name': name, 'idx': j,
                'code_row': code_row,
                'chern_label': code_row + 1,
                'chern_data': [code_row + 2, code_row + 3, code_row + 4],
                'chist_label': code_row + 5,
                'chist_data': chist_data,
                'zapol_row': zapol_row, 'blank': blank,
                'has_zapol': has_zapol,
            })
            row = blank + 1
        layout['korpus_sections'].append(section)
    layout['total_rows'] = row - 1
    return layout


# ─── СПРАВОЧНИК ROW LAYOUT ───────────────────────────────────────────────────
def calc_sprav_layout():
    layout = {}
    layout['parking_header'] = 3
    layout['parking_subheader'] = 4
    layout['parking_data_start'] = 5
    layout['parking_data_end'] = 10
    layout['parking_avg'] = 11

    layout['korpus'] = []
    for k in range(1, NUM_KORPUS + 1):
        header = 12 + (k - 1) * 10
        layout['korpus'].append({
            'k': k,
            'header': header,
            'subheader': header + 1,
            'data_start': header + 2,
            'data_end': header + 8,
            'avg': header + 9,
        })
    layout['total_rows'] = layout['korpus'][-1]['avg']
    return layout


# ─── COLUMN MAPPINGS ──────────────────────────────────────────────────────────
def chern_col_mapping(n_codes):
    result = []
    col = 2
    for _ in range(n_codes):
        result.append((col_letter(col), col_letter(col + 1), col_letter(col + 2)))
        col += 3
    return result


def chist_col_mapping(n_codes):
    result = []
    col = 2
    for _ in range(n_codes):
        result.append((col_letter(col), col_letter(col + 1), col_letter(col + 2),
                        col_letter(col + 3), col_letter(col + 4)))
        col += 5
    return result


def chern_itogo_start(n_codes):
    return 2 + n_codes * 3 + 1  # skip separator


def chist_itogo_start(n_codes):
    return 2 + n_codes * 5 + 1  # skip separator


# ─── HELPER ───────────────────────────────────────────────────────────────────
def style_cell(ws, row, col, value=None, font=None, fill=None, alignment=None, number_format=None, border=None):
    cell = ws.cell(row=row, column=col)
    if value is not None:
        cell.value = value
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if alignment:
        cell.alignment = alignment
    if number_format:
        cell.number_format = number_format
    if border:
        cell.border = border
    return cell


def merge_and_style(ws, min_row, min_col, max_row, max_col, value=None, font=None, fill=None, alignment=None):
    ws.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)
    cell = ws.cell(row=min_row, column=min_col)
    if value is not None:
        cell.value = value
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if alignment:
        cell.alignment = alignment


# ─── GENERATE СПРАВОЧНИК ─────────────────────────────────────────────────────
def gen_spravochnik(wb, otchet_layout, sprav_layout):
    ws = wb.create_sheet('СПРАВОЧНИК')

    ws['A1'] = 'Дата старта работ'
    ws['B1'] = START_DATE
    ws['B1'].number_format = 'mm-dd-yy'

    merge_and_style(ws, 2, 1, 2, 7, 'СПРАВОЧНИК КАТЕГОРИЙ ОТДЕЛКИ', FONT_BOLD)

    merge_and_style(ws, 3, 1, 3, 7, 'ПАРКИНГ', FONT_WHITE_BOLD_14, FILL_DARK_BLUE, ALIGN_CENTER)

    korpus_start_col = 9
    korpus_end_col = korpus_start_col + NUM_KORPUS
    merge_and_style(ws, 3, korpus_start_col, 3, korpus_end_col,
                    'Развесовка корпусов', FONT_WHITE_BOLD_14, FILL_DARK_BLUE, ALIGN_CENTER)

    obj_start_col = korpus_end_col + 2
    obj_end_col = obj_start_col + NUM_KORPUS + 1
    merge_and_style(ws, 3, obj_start_col, 3, obj_end_col,
                    'Развесовка объектов', FONT_WHITE_BOLD_14, FILL_DARK_BLUE, ALIGN_CENTER)

    headers_main = ['Код', 'Часть', 'Название', 'План м²', 'Вес черн.', 'Вес чист.', '∑ весов']
    for i, h in enumerate(headers_main, 1):
        style_cell(ws, 4, i, h, FONT_BOLD, FILL_LIGHT_BLUE, ALIGN_CENTER)

    for k in range(NUM_KORPUS):
        style_cell(ws, 4, korpus_start_col + k, f'К{k + 1}', FONT_BOLD, FILL_LIGHT_BLUE, ALIGN_CENTER)
    style_cell(ws, 4, korpus_end_col, 'Сумма', FONT_BOLD, FILL_LIGHT_BLUE, ALIGN_CENTER)

    style_cell(ws, 4, obj_start_col, 'Паркинг', FONT_BOLD, FILL_LIGHT_BLUE, ALIGN_CENTER)
    for k in range(NUM_KORPUS):
        style_cell(ws, 4, obj_start_col + 1 + k, f'К{k + 1}', FONT_BOLD, FILL_LIGHT_BLUE, ALIGN_CENTER)
    style_cell(ws, 4, obj_end_col, 'Сумма', FONT_BOLD, FILL_LIGHT_BLUE, ALIGN_CENTER)

    # Underground data rows 5-10
    for i, (code, name) in enumerate(UNDERGROUND_CODES):
        r = 5 + i
        style_cell(ws, r, 1, code, FONT_NORMAL, FILL_LIGHTEST_BLUE)
        style_cell(ws, r, 2, 'Подзем.', FONT_NORMAL, FILL_LIGHTEST_BLUE)
        style_cell(ws, r, 3, name, FONT_NORMAL, FILL_LIGHTEST_BLUE)
        style_cell(ws, r, 4, 0, FONT_NORMAL, FILL_LIGHTEST_BLUE,
                   number_format='_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-')
        style_cell(ws, r, 5, WEIGHT_CHERN, FONT_NORMAL, FILL_LIGHTEST_BLUE)
        style_cell(ws, r, 6, WEIGHT_CHIST, FONT_NORMAL, FILL_LIGHTEST_BLUE)
        style_cell(ws, r, 7, f'=E{r}+F{r}', FONT_NORMAL, FILL_LIGHTEST_BLUE)

    # Корпус weights (row 5)
    equal_weight = round(1 / NUM_KORPUS, 4)
    for k in range(NUM_KORPUS):
        style_cell(ws, 5, korpus_start_col + k, equal_weight, FONT_NORMAL, FILL_LIGHTEST_BLUE)
    kcl_s = col_letter(korpus_start_col)
    kcl_e = col_letter(korpus_start_col + NUM_KORPUS - 1)
    style_cell(ws, 5, korpus_end_col, f'=SUM({kcl_s}5:{kcl_e}5)', FONT_NORMAL, FILL_LIGHTEST_BLUE)

    # Object weights
    parking_weight = round(1 / (NUM_KORPUS + 1), 4)
    korpus_obj_weight = round((1 - parking_weight) / NUM_KORPUS, 4)
    style_cell(ws, 5, obj_start_col, parking_weight, FONT_NORMAL, FILL_LIGHTEST_BLUE)
    for k in range(NUM_KORPUS):
        style_cell(ws, 5, obj_start_col + 1 + k, korpus_obj_weight, FONT_NORMAL, FILL_LIGHTEST_BLUE)
    ocl_s = col_letter(obj_start_col)
    ocl_e = col_letter(obj_end_col - 1)
    style_cell(ws, 5, obj_end_col, f'=SUM({ocl_s}5:{ocl_e}5)', FONT_NORMAL, FILL_LIGHTEST_BLUE)

    # Row 11: Parking avg
    style_cell(ws, 11, 1, 'Паркинг', FONT_BOLD, FILL_LIGHTEST_BLUE)
    style_cell(ws, 11, 5, '=AVERAGE(E5:E10)', FONT_NORMAL, FILL_LIGHTEST_BLUE)
    style_cell(ws, 11, 6, '=AVERAGE(F5:F10)', FONT_NORMAL, FILL_LIGHTEST_BLUE)
    style_cell(ws, 11, 7, '=E11+F11', FONT_NORMAL, FILL_LIGHTEST_BLUE)

    # Корпус sections
    for ks in sprav_layout['korpus']:
        k = ks['k']
        hr = ks['header']
        merge_and_style(ws, hr, 1, hr, 7, f'КОРПУС {k}', FONT_WHITE_BOLD_14, FILL_DARK_BLUE, ALIGN_CENTER)

        sr = ks['subheader']
        for i, h in enumerate(headers_main, 1):
            style_cell(ws, sr, i, h, FONT_BOLD, FILL_LIGHT_BLUE, ALIGN_CENTER)

        for j, (code, name) in enumerate(ABOVE_GROUND_CODES):
            r = ks['data_start'] + j
            style_cell(ws, r, 1, code, FONT_NORMAL, FILL_LIGHTEST_BLUE)
            style_cell(ws, r, 2, 'Надзем.', FONT_NORMAL, FILL_LIGHTEST_BLUE)
            style_cell(ws, r, 3, name, FONT_NORMAL, FILL_LIGHTEST_BLUE)
            style_cell(ws, r, 4, 0, FONT_NORMAL, FILL_LIGHTEST_BLUE,
                       number_format='_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-')
            style_cell(ws, r, 5, WEIGHT_CHERN, FONT_NORMAL, FILL_LIGHTEST_BLUE)
            style_cell(ws, r, 6, WEIGHT_CHIST, FONT_NORMAL, FILL_LIGHTEST_BLUE)
            style_cell(ws, r, 7, f'=E{r}+F{r}', FONT_NORMAL, FILL_LIGHTEST_BLUE)

        ar = ks['avg']
        ds = ks['data_start']
        de = ks['data_end']
        style_cell(ws, ar, 1, f'Корпус {k}', FONT_BOLD, FILL_LIGHTEST_BLUE)
        style_cell(ws, ar, 5, f'=AVERAGE(E{ds}:E{de})', FONT_NORMAL, FILL_LIGHTEST_BLUE)
        style_cell(ws, ar, 6, f'=AVERAGE(F{ds}:F{de})', FONT_NORMAL, FILL_LIGHTEST_BLUE)
        style_cell(ws, ar, 7, f'=E{ar}+F{ar}', FONT_NORMAL, FILL_LIGHTEST_BLUE)

    ws.column_dimensions['A'].width = 20.14
    ws.column_dimensions['B'].width = 15.43
    ws.column_dimensions['C'].width = 38.57
    for c in range(4, obj_end_col + 2):
        ws.column_dimensions[col_letter(c)].width = 13.0

    return ws


# ─── GENERATE ПЛАН/ФАКТ ЧЕРН SHEET ───────────────────────────────────────────
def gen_chern_sheet(wb, sheet_name, codes, is_underground):
    ws = wb.create_sheet(sheet_name)
    n_codes = len(codes)
    last_data_col = 1 + n_codes * 3
    sep_col = last_data_col + 1
    itogo_col = sep_col + 1
    itogo_end = itogo_col + 2

    cm = chern_col_mapping(n_codes)

    merge_and_style(ws, 1, 1, 1, last_data_col, sheet_name, FONT_WHITE_BOLD, FILL_DARK_BLUE, ALIGN_CENTER)
    itogo_label = 'Итого подзем.' if is_underground else 'Итого надзем.'
    merge_and_style(ws, 1, itogo_col, 1, itogo_end, itogo_label, FONT_WHITE_BOLD, FILL_DARK_BLUE, ALIGN_CENTER)

    merge_and_style(ws, 2, 1, 3, 1, 'Дата', FONT_BOLD, FILL_LIGHT_BLUE, ALIGN_CENTER)
    area_label = 'Подземная часть' if is_underground else 'Надземная часть'
    merge_and_style(ws, 2, 2, 2, last_data_col, area_label, FONT_BOLD, FILL_LIGHTEST_BLUE, ALIGN_CENTER)

    col = 2
    for code, name in codes:
        merge_and_style(ws, 3, col, 3, col + 2, f'{code} {name}', FONT_BOLD, FILL_LIGHTEST_BLUE, ALIGN_CENTER)
        col += 3

    sub_headers = ['Пол', 'Стены', 'Потолок']
    col = 2
    for _ in codes:
        for sh in sub_headers:
            style_cell(ws, 4, col, sh, FONT_BOLD, FILL_LIGHTEST_BLUE, ALIGN_CENTER)
            col += 1

    # Итого headers row 2-4
    merge_and_style(ws, 2, itogo_col, 3, itogo_end, 'Итого', FONT_BOLD, FILL_YELLOW, ALIGN_CENTER)
    for si, sh in enumerate(sub_headers):
        style_cell(ws, 4, itogo_col + si, sh, FONT_BOLD, FILL_YELLOW, ALIGN_CENTER)

    for dr in range(DATA_ROWS):
        r = 5 + dr
        if dr == 0:
            style_cell(ws, r, 1, "=СПРАВОЧНИК!$B$1", number_format='dd\\.mm\\.yyyy')
        else:
            style_cell(ws, r, 1, f'=A{r - 1}+7', number_format='dd\\.mm\\.yyyy')

        col = 2
        for _ in codes:
            for _ in range(3):
                style_cell(ws, r, col, 0, number_format='#,##0.00')
                col += 1

        for si in range(3):
            src_cols = [cm[i][si] for i in range(n_codes)]
            formula = '+'.join([f'{c}{r}' for c in src_cols])
            style_cell(ws, r, itogo_col + si, f'={formula}', fill=FILL_YELLOW, number_format='#,##0.00')

    ws.column_dimensions['A'].width = 12.0
    for c in range(2, itogo_end + 1):
        ws.column_dimensions[col_letter(c)].width = 13.0
    ws.column_dimensions[col_letter(sep_col)].width = 3.0

    return ws


# ─── GENERATE ПЛАН/ФАКТ ЧИСТ SHEET ───────────────────────────────────────────
def gen_chist_sheet(wb, sheet_name, codes, is_underground):
    ws = wb.create_sheet(sheet_name)
    n_codes = len(codes)
    last_data_col = 1 + n_codes * 5
    sep_col = last_data_col + 1
    itogo_col = sep_col + 1
    itogo_end = itogo_col + 4

    cm = chist_col_mapping(n_codes)

    merge_and_style(ws, 1, 1, 1, itogo_end, sheet_name, FONT_WHITE_BOLD, FILL_BLUE, ALIGN_CENTER)

    merge_and_style(ws, 2, 1, 4, 1, 'Дата', FONT_WHITE_BOLD, FILL_BLUE, ALIGN_CENTER)
    area_label = 'Подземная часть' if is_underground else 'Надземная часть'
    merge_and_style(ws, 2, 2, 2, last_data_col, area_label, FONT_WHITE_BOLD, FILL_BLUE, ALIGN_CENTER)

    col = 2
    for code, name in codes:
        merge_and_style(ws, 3, col, 3, col + 4, f'{code} {name}', FONT_WHITE_BOLD, FILL_BLUE, ALIGN_CENTER)
        col += 5

    merge_and_style(ws, 3, itogo_col, 3, itogo_end, 'Итого', FONT_WHITE_BOLD, FILL_BLUE, ALIGN_CENTER)

    sub_headers = ['Пол', 'Стены', 'Потолок', 'Запол.', 'Мебл.']
    col = 2
    for _ in codes:
        for sh in sub_headers:
            style_cell(ws, 4, col, sh, FONT_WHITE_BOLD, FILL_BLUE, ALIGN_CENTER)
            col += 1

    for si, sh in enumerate(sub_headers):
        style_cell(ws, 4, itogo_col + si, sh, FONT_WHITE_BOLD, FILL_BLUE, ALIGN_CENTER)

    for dr in range(DATA_ROWS):
        r = 5 + dr
        if dr == 0:
            style_cell(ws, r, 1, "=СПРАВОЧНИК!$B$1", number_format='dd\\.mm\\.yyyy')
        else:
            style_cell(ws, r, 1, f'=A{r - 1}+7', number_format='dd\\.mm\\.yyyy')

        col = 2
        for _ in codes:
            for _ in range(5):
                style_cell(ws, r, col, 0, number_format='#,##0.00')
                col += 1

        for si in range(5):
            src_cols = [cm[i][si] for i in range(n_codes)]
            formula = '+'.join([f'{c}{r}' for c in src_cols])
            style_cell(ws, r, itogo_col + si, f'={formula}', number_format='#,##0.00')

    ws.column_dimensions['A'].width = 14.0
    for c in range(2, itogo_end + 1):
        ws.column_dimensions[col_letter(c)].width = 13.0
    ws.column_dimensions[col_letter(sep_col)].width = 10.0

    return ws


# ─── GENERATE ОТЧЕТ ──────────────────────────────────────────────────────────
def gen_otchet(wb, otchet_layout, sprav_layout):
    ws = wb.create_sheet('ОТЧЕТ')

    merge_and_style(ws, 1, 1, 1, 7, 'ОТЧЁТ ПО ОТДЕЛОЧНЫМ РАБОТАМ', FONT_WHITE_BOLD_14, FILL_BLUE, ALIGN_CENTER)
    merge_and_style(ws, 2, 1, 2, 7, '⚠ Автоматический. Не вносите данные напрямую.',
                    FONT_RED_BOLD, FILL_RED_WARN, ALIGN_CENTER)

    merge_and_style(ws, 4, 1, 4, 7, 'ПОДЗЕМНАЯ ЧАСТЬ', FONT_WHITE_BOLD, FILL_PURPLE, ALIGN_CENTER)

    col_headers = ['Код', 'Название', 'ВСЕГО м²', 'План на сегодня м²',
                   'Факт на сегодня м²', 'Дельта', 'Факт %']
    for i, h in enumerate(col_headers, 1):
        style_cell(ws, 5, i, h, FONT_WHITE_BOLD, FILL_BLUE, ALIGN_CENTER_WRAP)
    ws.row_dimensions[5].height = 30

    ws.column_dimensions['A'].width = 8.0
    ws.column_dimensions['B'].width = 36.0
    ws.column_dimensions['C'].width = 12.0
    ws.column_dimensions['D'].width = 13.0
    ws.column_dimensions['E'].width = 13.0
    ws.column_dimensions['F'].width = 12.0
    ws.column_dimensions['G'].width = 13.0

    u_chern_cm = chern_col_mapping(len(UNDERGROUND_CODES))
    u_chist_cm = chist_col_mapping(len(UNDERGROUND_CODES))

    def write_data_row(r, plan_sheet, fact_sheet, plan_col, fact_col, label):
        style_cell(ws, r, 2, f'    {label}', FONT_NORMAL, FILL_GRAY, ALIGN_LEFT_CENTER)
        style_cell(ws, r, 3, f"=SUM('{plan_sheet}'!{plan_col}5:{plan_col}90)",
                   FONT_NORMAL, FILL_GRAY, ALIGN_CENTER, '#,##0.00')
        style_cell(ws, r, 4,
                   f"=IFERROR(SUM('{plan_sheet}'!{plan_col}5:INDEX('{plan_sheet}'!{plan_col}5:{plan_col}90,"
                   f"MATCH(TODAY(),'{plan_sheet}'!A5:A90,1))),0)",
                   FONT_NORMAL, FILL_GRAY, ALIGN_CENTER, '#,##0.00')
        style_cell(ws, r, 5,
                   f"=IFERROR(SUM('{fact_sheet}'!{fact_col}5:INDEX('{fact_sheet}'!{fact_col}5:{fact_col}90,"
                   f"MATCH(TODAY(),'{fact_sheet}'!A5:A90,1))),0)",
                   FONT_NORMAL, FILL_GRAY, ALIGN_CENTER, '#,##0.00')
        style_cell(ws, r, 6, f'=E{r}-D{r}', FONT_NORMAL, FILL_GRAY, ALIGN_CENTER, '#,##0.00')
        style_cell(ws, r, 7, f'=IFERROR(E{r}/C{r},"0")', FONT_NORMAL, FILL_GRAY, ALIGN_CENTER, '0.0%')

    # Underground codes
    for u in otchet_layout['underground']:
        i = u['idx']
        cr = u['code_row']
        sprav_row = 5 + i

        style_cell(ws, cr, 1, u['code'], FONT_BOLD, FILL_GRAY, ALIGN_CENTER)
        style_cell(ws, cr, 2, u['name'], FONT_BOLD, FILL_GRAY, ALIGN_LEFT_CENTER)
        for c in range(3, 7):
            style_cell(ws, cr, c, None, FONT_NORMAL, FILL_GRAY, number_format='#,##0.00')

        chern_s = u['chern_data'][0]
        chern_e = u['chern_data'][-1]
        chist_s = u['chist_data'][0]
        chist_e = u['chist_data'][-1]
        g_formula = (f"=(SUM(ОТЧЕТ!E{chern_s}:E{chern_e})/SUM(ОТЧЕТ!C{chern_s}:C{chern_e})"
                     f"*СПРАВОЧНИК!E{sprav_row})+"
                     f"(SUM(ОТЧЕТ!E{chist_s}:E{chist_e})/SUM(ОТЧЕТ!C{chist_s}:C{chist_e})"
                     f"*СПРАВОЧНИК!F{sprav_row})")
        style_cell(ws, cr, 7, g_formula, FONT_NORMAL, FILL_GRAY, ALIGN_CENTER, '0.0%')

        # Черновая label + subtotal
        style_cell(ws, u['chern_label'], 2, '  Черновая', FONT_BOLD)
        chern_g = f"=SUM(E{u['chern_data'][0]}:E{u['chern_data'][-1]})/SUM(C{u['chern_data'][0]}:C{u['chern_data'][-1]})"
        style_cell(ws, u['chern_label'], 7, chern_g, number_format='0.0%')

        labels = ['Пол', 'Стены', 'Потолок']
        for di, dr in enumerate(u['chern_data']):
            write_data_row(dr, 'ПЛАН ЧЕРН ПОДЗЕМ', 'ФАКТ ЧЕРН ПОДЗЕМ', u_chern_cm[i][di], u_chern_cm[i][di], labels[di])

        # Чистовая label + subtotal
        style_cell(ws, u['chist_label'], 2, '  Чистовая', FONT_BOLD)
        chist_g = f"=SUM(E{u['chist_data'][0]}:E{u['chist_data'][-1]})/SUM(C{u['chist_data'][0]}:C{u['chist_data'][-1]})"
        style_cell(ws, u['chist_label'], 7, chist_g, number_format='0.0%')

        for di, dr in enumerate(u['chist_data']):
            write_data_row(dr, 'ПЛАН ЧИСТ ПОДЗЕМ', 'ФАКТ ЧИСТ ПОДЗЕМ', u_chist_cm[i][di], u_chist_cm[i][di], labels[di])

    # Корпус sections
    k_chern_cm = chern_col_mapping(len(ABOVE_GROUND_CODES))
    k_chist_cm = chist_col_mapping(len(ABOVE_GROUND_CODES))

    for ks in otchet_layout['korpus_sections']:
        k = ks['k']
        sr = ks['section_row']
        hr = ks['header_row']

        merge_and_style(ws, sr, 1, sr, 7, f'НАДЗЕМНАЯ ЧАСТЬ — К{k}',
                        FONT_WHITE_BOLD, FILL_PURPLE, ALIGN_CENTER)

        for i, h in enumerate(col_headers, 1):
            style_cell(ws, hr, i, h, FONT_WHITE_BOLD, FILL_BLUE, ALIGN_CENTER_WRAP)
        ws.row_dimensions[hr].height = 30

        plan_chern = f'ПЛАН ЧЕРН К{k}'
        fact_chern = f'ФАКТ ЧЕРН К{k}'
        plan_chist = f'ПЛАН ЧИСТ К{k}'
        fact_chist = f'ФАКТ ЧИСТ К{k}'

        sprav_k = sprav_layout['korpus'][k - 1]

        for c_info in ks['codes']:
            j = c_info['idx']
            cr = c_info['code_row']
            sprav_row = sprav_k['data_start'] + j

            style_cell(ws, cr, 1, c_info['code'], FONT_BOLD, FILL_GRAY, ALIGN_CENTER)
            style_cell(ws, cr, 2, c_info['name'], FONT_BOLD, FILL_GRAY, ALIGN_LEFT_CENTER)
            for col in range(3, 7):
                style_cell(ws, cr, col, None, FONT_NORMAL, FILL_GRAY, number_format='#,##0.00')

            chern_s = c_info['chern_data'][0]
            chern_e = c_info['chern_data'][-1]
            chist_s = c_info['chist_data'][0]
            chist_e = c_info['chist_data'][-1]

            chist_e_g = c_info['zapol_row'] if c_info['has_zapol'] else chist_e

            g_formula = (f"=(SUM(ОТЧЕТ!E{chern_s}:E{chern_e})/SUM(ОТЧЕТ!C{chern_s}:C{chern_e})"
                         f"*СПРАВОЧНИК!E{sprav_row})+"
                         f"(SUM(ОТЧЕТ!E{chist_s}:E{chist_e_g})/SUM(ОТЧЕТ!C{chist_s}:C{chist_e_g})"
                         f"*СПРАВОЧНИК!F{sprav_row})")
            style_cell(ws, cr, 7, g_formula, FONT_NORMAL, FILL_GRAY, ALIGN_CENTER, '0.0%')

            # Черновая + subtotal
            style_cell(ws, c_info['chern_label'], 2, '  Черновая', FONT_BOLD)
            chern_g = f"=SUM(E{c_info['chern_data'][0]}:E{c_info['chern_data'][-1]})/SUM(C{c_info['chern_data'][0]}:C{c_info['chern_data'][-1]})"
            style_cell(ws, c_info['chern_label'], 7, chern_g, number_format='0.0%')

            labels = ['Пол', 'Стены', 'Потолок']
            for di, dr in enumerate(c_info['chern_data']):
                write_data_row(dr, plan_chern, fact_chern, k_chern_cm[j][di], k_chern_cm[j][di], labels[di])

            # Чистовая + subtotal
            style_cell(ws, c_info['chist_label'], 2, '  Чистовая', FONT_BOLD)
            chist_g = f"=SUM(E{c_info['chist_data'][0]}:E{c_info['chist_data'][-1]})/SUM(C{c_info['chist_data'][0]}:C{c_info['chist_data'][-1]})"
            style_cell(ws, c_info['chist_label'], 7, chist_g, number_format='0.0%')

            for di, dr in enumerate(c_info['chist_data']):
                write_data_row(dr, plan_chist, fact_chist, k_chist_cm[j][di], k_chist_cm[j][di], labels[di])

            # Заполн.
            if c_info['has_zapol']:
                zapol_col = k_chist_cm[j][3]
                write_data_row(c_info['zapol_row'], plan_chist, fact_chist, zapol_col, zapol_col, 'Заполн. проемов')

    return ws


# ─── GENERATE SIGNAL ─────────────────────────────────────────────────────────
def gen_signal(wb, sprav_layout):
    ws = wb.create_sheet('SIGNAL')

    n_underground = len(UNDERGROUND_CODES)
    n_above = len(ABOVE_GROUND_CODES)

    chern_podz_last = col_letter(1 + n_underground * 3)
    chern_k_itogo_col = chern_itogo_start(n_above)
    chist_podz_itogo_col = chist_itogo_start(n_underground)
    chist_k_itogo_col = chist_itogo_start(n_above)

    cp_it_s = col_letter(chist_podz_itogo_col)
    cp_it_e = col_letter(chist_podz_itogo_col + 2)
    ck_it_s = col_letter(chist_k_itogo_col)
    ck_it_e = col_letter(chist_k_itogo_col + 2)

    cp_zapol = col_letter(chist_podz_itogo_col + 3)
    ck_zapol = col_letter(chist_k_itogo_col + 3)

    ck_chern_s = col_letter(chern_k_itogo_col)
    ck_chern_e = col_letter(chern_k_itogo_col + 2)

    style_cell(ws, 1, 1, 'Дата отчета', FONT_BOLD, FILL_SIGNAL_DATE, ALIGN_CENTER)
    style_cell(ws, 2, 1, START_DATE, FONT_NORMAL, FILL_SIGNAL_DATE, number_format='dd\\.mm\\.yyyy')

    cards = []
    cards.append(('8.1 Подземная часть: Общий', 'общий_подзем'))
    cards.append(('8.1 Подземная часть: Черновая', 'черн_подзем'))
    cards.append(('8.1 Подземная часть: Чистовая', 'чист_подзем'))

    for k in range(1, NUM_KORPUS + 1):
        cards.append((f'8.2 Надземная К{k}: Общий', f'общий_к{k}'))
        cards.append((f'8.2 Надземная К{k}: Черновая', f'черн_к{k}'))
        cards.append((f'8.2 Надземная К{k}: Чистовая', f'чист_к{k}'))

    cards.append(('8.2 Надзем. (все): Общий', 'общий_надзем_все'))
    cards.append(('8.2 Надзем. (все): Черновая', 'черн_надзем_все'))
    cards.append(('8.2 Надзем. (все): Чистовая', 'чист_надзем_все'))

    cards.append(('8. Отделка общий: Общий', 'общий_все'))
    cards.append(('8. Отделка общий: Черновая', 'черн_все'))
    cards.append(('8. Отделка общий: Чистовая', 'чист_все'))

    cards.append(('8. Отделка общий: Заполн. Проемов', 'запол_все'))
    cards.append(('8.1 Подземная часть: Заполн. Проемов', 'запол_подзем'))
    cards.append(('8.2 Надзем. (все): Заполн. Проемов', 'запол_надзем_все'))
    for k in range(1, NUM_KORPUS + 1):
        cards.append((f'8.2 Надземная К{k}: Заполн. Проемов', f'запол_к{k}'))

    card_cols = {}
    for ci, (label, key) in enumerate(cards):
        cs = 2 + ci * 6
        card_cols[key] = (cs + 2, cs + 3, cs)

    def cc(k):
        pc, fc, _ = card_cols[k]
        return col_letter(pc), col_letter(fc)

    for ci, (label, key) in enumerate(cards):
        cs = 2 + ci * 6
        val_col = cs + 1
        plan_col = cs + 2
        fact_col = cs + 3
        pcl = col_letter(plan_col)
        fcl = col_letter(fact_col)
        vcl = col_letter(val_col)

        style_cell(ws, 1, cs, 'Заголовок', FONT_BOLD, FILL_SIGNAL_LABEL)
        style_cell(ws, 1, val_col, 'План-факт по объемам')
        style_cell(ws, 2, cs, 'Статус', FONT_BOLD, FILL_SIGNAL_LABEL)
        style_cell(ws, 2, val_col, 'true')
        style_cell(ws, 3, cs, 'Url изображения', FONT_BOLD, FILL_SIGNAL_LABEL)
        style_cell(ws, 4, cs, 'Тип', FONT_BOLD, FILL_SIGNAL_LABEL)
        style_cell(ws, 4, val_col, 'planFact2')
        style_cell(ws, 6, cs, 'Дата', FONT_BOLD, FILL_SIGNAL_LABEL)
        style_cell(ws, 6, val_col, '=$A$2', number_format='dd\\.mm\\.yyyy')
        style_cell(ws, 7, cs, 'Гистограмма', FONT_BOLD, FILL_SIGNAL_LABEL)
        style_cell(ws, 7, val_col, 'false')
        style_cell(ws, 8, cs, 'По месяцам', FONT_BOLD, FILL_SIGNAL_LABEL)
        style_cell(ws, 8, val_col, 'false')
        style_cell(ws, 9, cs, 'Всего', FONT_BOLD, FILL_SIGNAL_LABEL)
        style_cell(ws, 9, val_col, f'=SUM({pcl}14:{pcl}79)', number_format='#,##0.00')
        style_cell(ws, 10, cs, 'Тип', FONT_BOLD, FILL_SIGNAL_LABEL)
        style_cell(ws, 10, val_col, label)
        style_cell(ws, 11, cs, 'Ед. изм.', FONT_BOLD, FILL_SIGNAL_LABEL)
        style_cell(ws, 11, val_col, 'м2')
        style_cell(ws, 12, cs, 'Учитывать дату карточки', FONT_BOLD, FILL_SIGNAL_LABEL)
        style_cell(ws, 12, val_col, 'true')
        style_cell(ws, 13, cs, 'Данные', FONT_BOLD, FILL_SIGNAL_DATA_HDR)

        for sr in range(14, 80):
            pf_row = sr - 9

            if sr == 14:
                style_cell(ws, sr, val_col, '=СПРАВОЧНИК!$B$1', number_format='dd\\.mm\\.yyyy')
            else:
                style_cell(ws, sr, val_col, f'={vcl}{sr - 1}+7', number_format='dd\\.mm\\.yyyy')

            plan_f, fact_f = _get_signal_formula(
                key, sr, pf_row, cc, sprav_layout,
                cp_it_s, cp_it_e, ck_it_s, ck_it_e,
                cp_zapol, ck_zapol,
                chern_podz_last,
                ck_chern_s, ck_chern_e)
            style_cell(ws, sr, plan_col, plan_f, number_format='#,##0.00')
            style_cell(ws, sr, fact_col, fact_f, fill=FILL_SIGNAL_FACT, number_format='#,##0.00')

    ws.column_dimensions['A'].width = 11
    card_widths = [25, 10, 10, 10, 8, 8]
    for ci in range(len(cards)):
        cs = 2 + ci * 6
        for offset, w in enumerate(card_widths):
            ws.column_dimensions[col_letter(cs + offset)].width = w
    return ws


def _get_signal_formula(key, sr, pf_row, cc, sprav_layout,
                        cp_it_s, cp_it_e, ck_it_s, ck_it_e,
                        cp_zapol, ck_zapol,
                        chern_podz_last,
                        ck_chern_s, ck_chern_e):

    if key == 'черн_подзем':
        return (f"=SUM('ПЛАН ЧЕРН ПОДЗЕМ'!B{pf_row}:{chern_podz_last}{pf_row})",
                f"=SUM('ФАКТ ЧЕРН ПОДЗЕМ'!B{pf_row}:{chern_podz_last}{pf_row})")

    if key == 'чист_подзем':
        return (f"=SUM('ПЛАН ЧИСТ ПОДЗЕМ'!{cp_it_s}{pf_row}:{cp_it_e}{pf_row})",
                f"=SUM('ФАКТ ЧИСТ ПОДЗЕМ'!{cp_it_s}{pf_row}:{cp_it_e}{pf_row})")

    if key == 'общий_подзем':
        cp, cf = cc('черн_подзем')
        tp, tf = cc('чист_подзем')
        return (f"={cp}{sr}*СПРАВОЧНИК!$E$11+{tp}{sr}*СПРАВОЧНИК!$F$11",
                f"={cf}{sr}*СПРАВОЧНИК!$E$11+{tf}{sr}*СПРАВОЧНИК!$F$11")

    if key.startswith('черн_к') and key[6:].isdigit():
        k = int(key[6:])
        return (f"=SUM('ПЛАН ЧЕРН К{k}'!{ck_chern_s}{pf_row}:{ck_chern_e}{pf_row})",
                f"=SUM('ФАКТ ЧЕРН К{k}'!{ck_chern_s}{pf_row}:{ck_chern_e}{pf_row})")

    if key.startswith('чист_к') and key[6:].isdigit():
        k = int(key[6:])
        return (f"=SUM('ПЛАН ЧИСТ К{k}'!{ck_it_s}{pf_row}:{ck_it_e}{pf_row})",
                f"=SUM('ФАКТ ЧИСТ К{k}'!{ck_it_s}{pf_row}:{ck_it_e}{pf_row})")

    if key.startswith('общий_к') and key[7:].isdigit():
        k = int(key[7:])
        avg_row = sprav_layout['korpus'][k - 1]['avg']
        cp, cf = cc(f'черн_к{k}')
        tp, tf = cc(f'чист_к{k}')
        return (f"={cp}{sr}*СПРАВОЧНИК!$E${avg_row}+{tp}{sr}*СПРАВОЧНИК!$F${avg_row}",
                f"={cf}{sr}*СПРАВОЧНИК!$E${avg_row}+{tf}{sr}*СПРАВОЧНИК!$F${avg_row}")

    if key == 'черн_надзем_все':
        parts_p = '+'.join([cc(f'черн_к{k}')[0] + str(sr) for k in range(1, NUM_KORPUS + 1)])
        parts_f = '+'.join([cc(f'черн_к{k}')[1] + str(sr) for k in range(1, NUM_KORPUS + 1)])
        return f'={parts_p}', f'={parts_f}'

    if key == 'чист_надзем_все':
        parts_p = '+'.join([cc(f'чист_к{k}')[0] + str(sr) for k in range(1, NUM_KORPUS + 1)])
        parts_f = '+'.join([cc(f'чист_к{k}')[1] + str(sr) for k in range(1, NUM_KORPUS + 1)])
        return f'={parts_p}', f'={parts_f}'

    if key == 'общий_надзем_все':
        cp, cf = cc('черн_надзем_все')
        tp, tf = cc('чист_надзем_все')
        avg_rows = [sprav_layout['korpus'][k - 1]['avg'] for k in range(1, NUM_KORPUS + 1)]
        avg_e = '+'.join([f'СПРАВОЧНИК!$E${r}' for r in avg_rows])
        avg_f_w = '+'.join([f'СПРАВОЧНИК!$F${r}' for r in avg_rows])
        n = NUM_KORPUS
        return (f"={cp}{sr}*(({avg_e})/{n})+{tp}{sr}*(({avg_f_w})/{n})",
                f"={cf}{sr}*(({avg_e})/{n})+{tf}{sr}*(({avg_f_w})/{n})")

    if key == 'черн_все':
        np, nf = cc('черн_надзем_все')
        pp, pf_ = cc('черн_подзем')
        return f'={np}{sr}+{pp}{sr}', f'={nf}{sr}+{pf_}{sr}'

    if key == 'чист_все':
        np, nf = cc('чист_надзем_все')
        pp, pf_ = cc('чист_подзем')
        return f'={np}{sr}+{pp}{sr}', f'={nf}{sr}+{pf_}{sr}'

    if key == 'общий_все':
        cp, cf = cc('черн_все')
        tp, tf = cc('чист_все')
        avg_rows = [11] + [sprav_layout['korpus'][k - 1]['avg'] for k in range(1, NUM_KORPUS + 1)]
        avg_e = '+'.join([f'СПРАВОЧНИК!$E${r}' for r in avg_rows])
        avg_f_w = '+'.join([f'СПРАВОЧНИК!$F${r}' for r in avg_rows])
        n = len(avg_rows)
        return (f"={cp}{sr}*(({avg_e})/{n})+{tp}{sr}*(({avg_f_w})/{n})",
                f"={cf}{sr}*(({avg_e})/{n})+{tf}{sr}*(({avg_f_w})/{n})")

    if key == 'запол_подзем':
        return (f"='ПЛАН ЧИСТ ПОДЗЕМ'!{cp_zapol}{pf_row}",
                f"='ФАКТ ЧИСТ ПОДЗЕМ'!{cp_zapol}{pf_row}")

    if key.startswith('запол_к') and key[7:].isdigit():
        k = int(key[7:])
        return (f"='ПЛАН ЧИСТ К{k}'!{ck_zapol}{pf_row}",
                f"='ФАКТ ЧИСТ К{k}'!{ck_zapol}{pf_row}")

    if key == 'запол_надзем_все':
        parts_p = '+'.join([cc(f'запол_к{k}')[0] + str(sr) for k in range(1, NUM_KORPUS + 1)])
        parts_f = '+'.join([cc(f'запол_к{k}')[1] + str(sr) for k in range(1, NUM_KORPUS + 1)])
        return f'={parts_p}', f'={parts_f}'

    if key == 'запол_все':
        np, nf = cc('запол_надзем_все')
        pp, pf_ = cc('запол_подзем')
        return f'={np}{sr}+{pp}{sr}', f'={nf}{sr}+{pf_}{sr}'

    return '=0', '=0'


# ─── MAIN ────────────────────────────────────────────────────────────────────
def main():
    print("Calculating layouts...")
    otchet_layout = calc_otchet_layout()
    sprav_layout = calc_sprav_layout()

    print(f"ОТЧЕТ: {otchet_layout['total_rows']} rows")
    print(f"СПРАВОЧНИК: {sprav_layout['total_rows']} rows")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    print("Generating СПРАВОЧНИК...")
    gen_spravochnik(wb, otchet_layout, sprav_layout)

    print("Generating SIGNAL...")
    gen_signal(wb, sprav_layout)

    print("Generating ОТЧЕТ...")
    gen_otchet(wb, otchet_layout, sprav_layout)

    print("Generating ПЛАН/ФАКТ ПОДЗЕМ...")
    gen_chern_sheet(wb, 'ПЛАН ЧЕРН ПОДЗЕМ', UNDERGROUND_CODES, True)
    gen_chern_sheet(wb, 'ФАКТ ЧЕРН ПОДЗЕМ', UNDERGROUND_CODES, True)
    gen_chist_sheet(wb, 'ПЛАН ЧИСТ ПОДЗЕМ', UNDERGROUND_CODES, True)
    gen_chist_sheet(wb, 'ФАКТ ЧИСТ ПОДЗЕМ', UNDERGROUND_CODES, True)

    for k in range(1, NUM_KORPUS + 1):
        print(f"Generating ПЛАН/ФАКТ К{k}...")
        gen_chern_sheet(wb, f'ПЛАН ЧЕРН К{k}', ABOVE_GROUND_CODES, False)
        gen_chern_sheet(wb, f'ФАКТ ЧЕРН К{k}', ABOVE_GROUND_CODES, False)
        gen_chist_sheet(wb, f'ПЛАН ЧИСТ К{k}', ABOVE_GROUND_CODES, False)
        gen_chist_sheet(wb, f'ФАКТ ЧИСТ К{k}', ABOVE_GROUND_CODES, False)

    # Tab colors
    tab_colors = {
        'СПРАВОЧНИК': '2F5496',
        'SIGNAL': '002060',
        'ОТЧЕТ': '002060',
        'ПЛАН ЧЕРН ПОДЗЕМ': '00B0F0',
        'ПЛАН ЧИСТ ПОДЗЕМ': '00B0F0',
        'ФАКТ ЧЕРН ПОДЗЕМ': '92D050',
        'ФАКТ ЧИСТ ПОДЗЕМ': '92D050',
    }
    for k in range(1, NUM_KORPUS + 1):
        tab_colors[f'ПЛАН ЧЕРН К{k}'] = '0070C0'
        tab_colors[f'ПЛАН ЧИСТ К{k}'] = '0070C0'
        tab_colors[f'ФАКТ ЧЕРН К{k}'] = '00B050'
        tab_colors[f'ФАКТ ЧИСТ К{k}'] = '00B050'
    for name, color in tab_colors.items():
        if name in wb.sheetnames:
            wb[name].sheet_properties.tabColor = color

    sheet_names = wb.sheetnames
    print(f"Total sheets: {len(sheet_names)}")
    for s in sheet_names:
        print(f"  - {s}")

    print(f"\nSaving to {OUTPUT}...")
    wb.save(OUTPUT)
    print("Done!")


if __name__ == '__main__':
    main()
